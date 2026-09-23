import os
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

def create_professional_docx(md_path, docx_path):
    doc = Document()
    
    # Set default style to Times New Roman, Size 12
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(12)
    
    # helper for spacing
    def set_spacing(paragraph):
        paragraph_format = paragraph.paragraph_format
        paragraph_format.line_spacing = 1.5
        paragraph_format.space_after = Pt(10)

    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    print(f"--- Converting {md_path} to {docx_path} ---")

    for line in lines:
        line = line.strip()
        if not line:
            continue

        # Page Breaks for Chapters
        if line.startswith('## CHAPTER') or line.startswith('### 📑') or line.startswith('## 📜 1.'):
            doc.add_page_break()

        if line.startswith('# '):
            # level 1 heading (Main Title)
            p = doc.add_heading(line[2:], level=0)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif line.startswith('## '):
            # level 2 heading (Chapters)
            p = doc.add_heading(line[3:], level=1)
            set_spacing(p)
        elif line.startswith('### '):
            # level 3 heading (Sub-points)
            p = doc.add_heading(line[4:], level=2)
            set_spacing(p)
        elif line.startswith('- '):
            # Bullet point
            p = doc.add_paragraph(line[2:], style='List Bullet')
            set_spacing(p)
        elif line.startswith('**') and line.endswith('**'):
            # Bold stand-alone line
            p = doc.add_paragraph()
            run = p.add_run(line.strip('*'))
            run.bold = True
            set_spacing(p)
        elif '---' in line:
            # Horizontal rule -> Page break
            doc.add_page_break()
        else:
            # Normal paragraph
            # Basic support for inline bolding
            p = doc.add_paragraph()
            parts = line.split('**')
            for i, part in enumerate(parts):
                run = p.add_run(part)
                if i % 2 != 0:
                    run.bold = True
            set_spacing(p)

    # Force all text to Times New Roman (Surgical fix for python-docx quirk)
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            run.font.name = 'Times New Roman'
            rPr = run._element.get_or_add_rPr()
            rFonts = rPr.get_or_add_rFonts()
            rFonts.set(qn('w:ascii'), 'Times New Roman')
            rFonts.set(qn('w:hAnsi'), 'Times New Roman')

    doc.save(docx_path)
    print(f"--- SUCCESS: {docx_path} generated with Professional Academic Formatting ---")

if __name__ == "__main__":
    MD_INPUT = r"c:\Users\sreya\Desktop\Project\documentation_docs\FINAL_REPORT.md"
    DOCX_OUTPUT = r"c:\Users\sreya\Desktop\Project\documentation_docs\CyberShield_Final_Academic_Report.docx"
    
    if os.path.exists(MD_INPUT):
        create_professional_docx(MD_INPUT, DOCX_OUTPUT)
    else:
        print(f"Error: Could not find {MD_INPUT}")
